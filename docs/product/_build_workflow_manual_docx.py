# -*- coding: utf-8 -*-
"""Build OPS业务操作手册-IP组与工作任务.docx from Markdown + workflow screenshots."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).resolve().parent
MD_PATH = ROOT / "OPS业务操作手册-IP组与工作任务.md"
SHOT_DIR = ROOT / "manual-screenshots" / "workflow"
OUT_DOCX = ROOT / "OPS业务操作手册-IP组与工作任务.docx"

IMG_RE = re.compile(r"^!\[([^\]]*)\]\(([^)]+)\)\s*$")
CAPTION_RE = re.compile(r"^\*(.+)\*\s*$")


def set_run_font(run, name="微软雅黑", size_pt=11, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size_pt)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_caption(doc: Document, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, size_pt=9, color=RGBColor(0x66, 0x66, 0x66))
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(12)


def resolve_image(rel_path: str) -> Path:
    rel = rel_path.replace("./", "").replace("\\", "/")
    if rel.startswith("manual-screenshots/"):
        return ROOT / rel
    return SHOT_DIR / Path(rel).name


def add_image(doc: Document, path: Path, caption: str, max_width_in=6.2):
    if not path.exists() or path.stat().st_size < 5000:
        note = doc.add_paragraph()
        run = note.add_run(f"[截图缺失或无效: {path.name}]")
        set_run_font(run, size_pt=9, color=RGBColor(0xCC, 0x00, 0x00))
        return
    with Image.open(path) as im:
        w, h = im.size
    width_in = max_width_in
    height_in = width_in * (h / w)
    if height_in > 7.5:
        height_in = 7.5
        width_in = height_in * (w / h)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width_in))
    add_caption(doc, caption)


def strip_md_inline(text: str) -> str:
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    return text.replace("**", "").replace("__", "").replace("`", "")


def add_paragraph_with_inline(doc: Document, text: str, style=None):
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = p.add_run(part[2:-2])
            set_run_font(run, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            run = p.add_run(part[1:-1])
            set_run_font(run, name="Consolas", size_pt=10)
            run.font.color.rgb = RGBColor(0x33, 0x33, 0x99)
        else:
            cleaned = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", part)
            run = p.add_run(cleaned)
            set_run_font(run)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    return p


def add_table(doc: Document, rows: list[list[str]]):
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j in range(cols):
            cell = table.rows[i].cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            val = row[j] if j < len(row) else ""
            run = p.add_run(strip_md_inline(val))
            set_run_font(run, size_pt=9, bold=(i == 0))
            if i == 0:
                shading = cell._element.get_or_add_tcPr()
                shd = OxmlElement("w:shd")
                shd.set(qn("w:fill"), "E7EEF7")
                shd.set(qn("w:val"), "clear")
                shading.append(shd)
    doc.add_paragraph()


def convert():
    md = MD_PATH.read_text(encoding="utf-8")
    lines = md.splitlines()

    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)

    styles = doc.styles
    styles["Normal"].font.name = "微软雅黑"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    styles["Normal"].font.size = Pt(11)

    i = 0
    in_code = False
    code_buf: list[str] = []
    table_buf: list[list[str]] = []

    def flush_table():
        nonlocal table_buf
        if table_buf:
            cleaned = []
            for row in table_buf:
                if all(re.fullmatch(r":?-+:?", c.strip()) for c in row if c.strip()):
                    continue
                cleaned.append(row)
            add_table(doc, cleaned)
            table_buf = []

    def flush_code():
        nonlocal code_buf
        if not code_buf:
            return
        p = doc.add_paragraph()
        run = p.add_run("\n".join(code_buf))
        set_run_font(run, name="Consolas", size_pt=9)
        run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
        p.paragraph_format.left_indent = Cm(0.4)
        p.paragraph_format.space_after = Pt(8)
        code_buf = []

    while i < len(lines):
        line = lines[i]

        if line.strip().startswith("```"):
            if in_code:
                flush_code()
                in_code = False
            else:
                flush_table()
                in_code = True
                code_buf = []
            i += 1
            continue

        if in_code:
            code_buf.append(line)
            i += 1
            continue

        img_match = IMG_RE.match(line.strip())
        if img_match:
            flush_table()
            alt, rel_path = img_match.group(1), img_match.group(2)
            caption = alt
            if i + 1 < len(lines):
                cap_match = CAPTION_RE.match(lines[i + 1].strip())
                if cap_match:
                    caption = cap_match.group(1)
                    i += 1
            add_image(doc, resolve_image(rel_path), caption)
            i += 1
            continue

        if line.strip().startswith("|"):
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            table_buf.append(cells)
            i += 1
            continue
        flush_table()

        if re.fullmatch(r"-{3,}", line.strip()):
            i += 1
            continue

        m = re.match(r"^(#{1,4})\s+(.*)$", line)
        if m:
            level = len(m.group(1))
            title_plain = strip_md_inline(m.group(2).strip())
            style = {1: "Heading 1", 2: "Heading 2", 3: "Heading 3", 4: "Heading 4"}.get(level, "Heading 4")
            p = doc.add_paragraph(title_plain, style=style)
            for run in p.runs:
                set_run_font(run, size_pt={1: 18, 2: 14, 3: 12, 4: 11}[level], bold=True)
            i += 1
            continue

        if line.startswith(">"):
            text = line.lstrip("> ").strip()
            p = add_paragraph_with_inline(doc, text)
            for run in p.runs:
                run.italic = True
                run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
            i += 1
            continue

        if re.match(r"^(\s*)[-*]\s+", line) or re.match(r"^(\s*)\d+\.\s+", line):
            text = re.sub(r"^(\s*)([-*]|\d+\.)\s+", "", line)
            p = add_paragraph_with_inline(doc, "• " + text)
            p.paragraph_format.left_indent = Cm(0.5)
            i += 1
            continue

        if not line.strip():
            i += 1
            continue

        add_paragraph_with_inline(doc, line)
        i += 1

    flush_table()
    flush_code()

    note = doc.add_paragraph()
    run = note.add_run(
        "\n—— 操作截图于本地联调环境（http://127.0.0.1:5777）Playwright 自动采集；"
        "演示数据为空属常态。账号口令以实际环境为准，勿将密钥写入手册。"
    )
    set_run_font(run, size_pt=9, color=RGBColor(0x88, 0x88, 0x88))

    doc.save(str(OUT_DOCX))
    print(f"Wrote {OUT_DOCX} ({OUT_DOCX.stat().st_size:,} bytes)")


if __name__ == "__main__":
    convert()
